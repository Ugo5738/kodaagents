import time
from io import BytesIO

from django.core.files.base import ContentFile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from koda.config.logging_config import configure_logger

logger = configure_logger(__name__)

def generate_resume_docx(improved_resume_dict, filename):
    start_time = time.time()

    doc = Document()

    # Set up the document
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Add contact information
    contact_info = improved_resume_dict.get('contact', {})
    doc.add_paragraph(contact_info.get('name', '')).alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Add job title
    job_title = contact_info.get('job_title', '')
    if job_title:
        doc.add_paragraph(job_title).alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Combine email, phone, and address
    contact_details = []
    if contact_info.get('email'):
        contact_details.append(contact_info['email'])
    if contact_info.get('phone'):
        contact_details.append(contact_info['phone'])
    if contact_info.get('address'):
        contact_details.append(contact_info['address'])

    doc.add_paragraph(' | '.join(contact_details)).alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Add LinkedIn URL
    linkedin_url = contact_info.get('linkedIn', '')
    if linkedin_url:
        doc.add_paragraph(linkedin_url).alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Add summary
    doc.add_heading('Summary', level=1)
    doc.add_paragraph(improved_resume_dict.get('summary', ''))

    # Add experiences
    doc.add_heading('Experience', level=1)
    for exp in improved_resume_dict.get('experiences', {}).values():
        p = doc.add_paragraph()
        p.add_run(f"{exp.get('company_name', '')} - {exp.get('job_role', '')}").bold = True
        p.add_run(f"\n{exp.get('start_date', '')} - {exp.get('end_date', '')}")

        # Add location
        location = exp.get('location', '')
        if location:
            p.add_run(f" | {location}")

        for desc in exp.get('job_description', []):
            doc.add_paragraph(desc, style='List Bullet')

    # Add education
    doc.add_heading('Education', level=1)
    education_data = improved_resume_dict.get('education', [])
    if isinstance(education_data, dict):
        education_data = [education_data]
    for edu in education_data:
        p = doc.add_paragraph()
        p.add_run(f"{edu.get('degree', '')} - {edu.get('institution', '')}").bold = True
        p.add_run(f"\n{edu.get('end_date', '')}")

        # Add location
        location = edu.get('location', '')
        if location:
            p.add_run(f" | {location}")

    # Add skills
    doc.add_heading('Skills', level=1)
    doc.add_paragraph(', '.join(improved_resume_dict.get('skills', [])))

    # Add certifications
    certifications = improved_resume_dict.get('certifications', [])
    if certifications:
        doc.add_heading('Certifications', level=1)
        for cert in certifications:
            p = doc.add_paragraph()
            p.add_run(f"{cert.get('title', '')} - {cert.get('issuing_organization', '')}").bold = True
            p.add_run(f"\n{cert.get('date_obtained', '')}")

    docx_buffer = BytesIO()
    doc.save(docx_buffer)
    docx_buffer.seek(0)

    total = time.time() - start_time
    logger.info(f"DOCX CREATION TIME: {total}")
    return ContentFile(docx_buffer.getvalue(), name=filename)


def generate_cover_letter_docx(cover_letter_content, filename):
    start_time = time.time()

    doc = Document()

    # Set up the document
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Add content
    for key, value in cover_letter_content.items():
        if key == 'recipient':
            doc.add_paragraph(value)
        elif key == 'body':
            paragraphs = value.split('\n\n')
            for para in paragraphs:
                doc.add_paragraph(para)
        elif key == 'closing':
            doc.add_paragraph(value)
        elif key == 'name':
            doc.add_paragraph(value)

    docx_buffer = BytesIO()
    doc.save(docx_buffer)
    docx_buffer.seek(0)

    total = time.time() - start_time
    logger.info(f"DOCX CREATION TIME: {total}")
    return ContentFile(docx_buffer.getvalue(), name=filename)
